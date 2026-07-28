import argparse
import os
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Sequence

from openpyxl import load_workbook


@dataclass
class ExcelBatch:
    table_name: str
    source_file: str
    source_sheet: str
    rows: List[Dict[str, Any]]


TABLE_COLUMNS: Mapping[str, Sequence[str]] = {
    "project_progress": (
        "project_id",
        "project_name",
        "customer",
        "product_type",
        "apqp_phase",
        "current_status",
        "planned_start_date",
        "planned_finish_date",
        "actual_finish_date",
        "owner",
        "delay_days",
        "progress_note",
    ),
    "project_risks": (
        "risk_id",
        "project_id",
        "project_name",
        "risk_category",
        "risk_description",
        "severity_s",
        "occurrence_o",
        "detection_d",
        "rpn",
        "risk_level",
        "mitigation_action",
        "owner",
        "planned_close_date",
        "closure_status",
    ),
    "project_issues": (
        "issue_id",
        "project_id",
        "project_name",
        "issue_source",
        "issue_description",
        "found_date",
        "responsible_department",
        "owner",
        "planned_close_date",
        "actual_close_date",
        "current_status",
        "resolution_note",
    ),
    "apqp_deliverables": (
        "deliverable_id",
        "project_id",
        "project_name",
        "apqp_phase",
        "deliverable_name",
        "is_completed",
        "review_status",
        "owner",
        "planned_submit_date",
        "actual_submit_date",
        "missing_or_reject_reason",
    ),
}


INT_COLUMNS = {"delay_days", "severity_s", "occurrence_o", "detection_d", "rpn"}
DATE_COLUMNS = {
    "planned_start_date",
    "planned_finish_date",
    "actual_finish_date",
    "planned_close_date",
    "actual_close_date",
    "found_date",
    "planned_submit_date",
    "actual_submit_date",
}
SOURCE_COLUMNS = ("source_file", "source_sheet")

HEADER_ALIASES = {
    "交付物id": "deliverable_id",
    "交付物_id": "deliverable_id",
    "项目id": "project_id",
    "项目_id": "project_id",
    "项目名称": "project_name",
    "顾客": "customer",
    "客户": "customer",
    "产品类型": "product_type",
    "当前状态": "current_status",
    "计划开始日期": "planned_start_date",
    "计划完成日期": "planned_finish_date",
    "实际完成日期": "actual_finish_date",
    "所有者": "owner",
    "负责人": "owner",
    "延迟天数": "delay_days",
    "延期天数": "delay_days",
    "进度说明": "progress_note",
    "风险标识": "risk_id",
    "风险编号": "risk_id",
    "风险类别": "risk_category",
    "风险描述": "risk_description",
    "严重程度": "severity_s",
    "发生概率": "occurrence_o",
    "探测度": "detection_d",
    "风险等级": "risk_level",
    "缓解措施": "mitigation_action",
    "应对措施": "mitigation_action",
    "计划关闭日期": "planned_close_date",
    "关闭状态": "closure_status",
    "问题编号": "issue_id",
    "问题来源": "issue_source",
    "问题描述": "issue_description",
    "发现日期": "found_date",
    "负责部门": "responsible_department",
    "责任部门": "responsible_department",
    "实际关闭日期": "actual_close_date",
    "决议说明": "resolution_note",
    "处理结论": "resolution_note",
    "交付物名称": "deliverable_name",
    "已完成": "is_completed",
    "是否完成": "is_completed",
    "审核状态": "review_status",
    "计划提交日期": "planned_submit_date",
    "实际提交日期": "actual_submit_date",
    "缺失或拒绝原因": "missing_or_reject_reason",
    "缺失或退回原因": "missing_or_reject_reason",
}


def extract_excel_batches(input_dir: str | Path) -> List[ExcelBatch]:
    root = Path(input_dir)
    batches: List[ExcelBatch] = []
    for path in sorted(root.rglob("*.xlsx")):
        batches.extend(_extract_workbook(path))
    return batches


def import_excel_directory(input_dir: str | Path, connection: Any, reset_tables: bool = True) -> Dict[str, int]:
    batches = extract_excel_batches(input_dir)
    touched_tables = sorted({batch.table_name for batch in batches})

    with connection.cursor() as cursor:
        for table_name in touched_tables:
            _create_table(cursor, table_name)
            if reset_tables:
                cursor.execute(f"TRUNCATE TABLE {table_name}")

        for batch in batches:
            if not batch.rows:
                continue
            _insert_batch(cursor, batch)

    connection.commit()
    return {
        "files": len({batch.source_file for batch in batches}),
        "sheets": len(batches),
        "rows": sum(len(batch.rows) for batch in batches),
    }


def connect_mysql_from_env() -> Any:
    try:
        import mysql.connector
    except ImportError as exc:
        raise RuntimeError(
            "mysql-connector-python is required. Install dependencies with: "
            "pip install -r back/requirements.txt"
        ) from exc

    return mysql.connector.connect(
        host=os.getenv("MYSQL_HOST", "127.0.0.1"),
        port=int(os.getenv("MYSQL_PORT", "3306")),
        user=os.getenv("MYSQL_USER", "root"),
        password=os.getenv("MYSQL_PASSWORD", ""),
        database=os.getenv("MYSQL_DATABASE", "apqp"),
        charset="utf8mb4",
    )


def _extract_workbook(path: Path) -> List[ExcelBatch]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    batches: List[ExcelBatch] = []

    for worksheet in workbook.worksheets:
        rows_iter = worksheet.iter_rows(values_only=True)
        raw_headers = next(rows_iter, None)
        if not raw_headers:
            continue

        headers = [_normalize_header(value) for value in raw_headers]
        table_name = _identify_table(headers)
        if table_name is None:
            continue

        expected_columns = TABLE_COLUMNS[table_name]
        data_rows: List[Dict[str, Any]] = []
        for raw_row in rows_iter:
            if not raw_row or all(value is None or str(value).strip() == "" for value in raw_row):
                continue
            row_by_header = {
                header: _normalize_value(value)
                for header, value in zip(headers, raw_row)
                if header
            }
            record = {column: row_by_header.get(column) for column in expected_columns}
            record["source_file"] = path.name
            record["source_sheet"] = worksheet.title
            data_rows.append(record)

        batches.append(
            ExcelBatch(
                table_name=table_name,
                source_file=path.name,
                source_sheet=worksheet.title,
                rows=data_rows,
            )
        )

    workbook.close()
    return batches


def _identify_table(headers: Sequence[str]) -> str | None:
    header_set = {header for header in headers if header}
    for table_name, columns in TABLE_COLUMNS.items():
        if set(columns).issubset(header_set):
            return table_name
    return None


def _create_table(cursor: Any, table_name: str) -> None:
    columns_sql = []
    for column in (*TABLE_COLUMNS[table_name], *SOURCE_COLUMNS):
        columns_sql.append(f"{column} {_mysql_type(column)}")

    cursor.execute(
        f"""
        CREATE TABLE IF NOT EXISTS {table_name} (
            id BIGINT AUTO_INCREMENT PRIMARY KEY,
            {", ".join(columns_sql)},
            ingested_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
        """
    )


def _insert_batch(cursor: Any, batch: ExcelBatch) -> None:
    columns = [*TABLE_COLUMNS[batch.table_name], *SOURCE_COLUMNS]
    placeholders = ", ".join(f"%({column})s" for column in columns)
    column_sql = ", ".join(columns)
    cursor.executemany(
        f"INSERT INTO {batch.table_name} ({column_sql}) VALUES ({placeholders})",
        batch.rows,
    )


def _mysql_type(column: str) -> str:
    if column in INT_COLUMNS:
        return "INT NULL"
    if column in DATE_COLUMNS:
        return "DATE NULL"
    if column.endswith("_description") or column.endswith("_note") or column.endswith("_action") or column.endswith("_reason"):
        return "TEXT NULL"
    return "VARCHAR(255) NULL"


def _normalize_header(value: Any) -> str:
    normalized = str(value or "").strip().lower().replace(" ", "_")
    compact = normalized.replace("_", "")
    return HEADER_ALIASES.get(normalized) or HEADER_ALIASES.get(compact) or normalized


def _normalize_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        return text if text else None
    return value


def main(argv: List[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Extract APQP Excel data and load it into MySQL.")
    parser.add_argument("--input-dir", default="knowledge_docs")
    parser.add_argument("--no-reset", action="store_true", help="Append rows instead of truncating known tables first.")
    args = parser.parse_args(argv)

    connection = connect_mysql_from_env()
    try:
        summary = import_excel_directory(args.input_dir, connection, reset_tables=not args.no_reset)
    finally:
        connection.close()

    print(f"Imported {summary['rows']} rows from {summary['files']} Excel files across {summary['sheets']} sheets.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())


