import re
import zipfile
import xml.etree.ElementTree as ET
from typing import Dict, List, Tuple, Any
import pandas as pd
import docx


class CrossDocumentValidator:
    PROCESS_COL_KEYWORDS = {
        "工序", "过程", "步骤", "工位", "操作", "流程", "工步", "制程", "加工", "工位名称",
        "过程步骤", "流程步骤", "生产步骤", "作业", "作业步骤", "工序号", "工序名称",
        "process", "step", "operation", "workstation", "procedure", "stage", "operation_name"
    }
    CHARACTERISTIC_COL_KEYWORDS = {
        "特性", "产品特性", "过程特性", "关键特性", "特殊特性", "尺寸", "规格", "参数",
        "产品特殊特性", "过程特殊特性", "重要特性", "特性值", "控制特性",
        "characteristic", "spec", "dimension", "feature", "product_char", "process_char"
    }
    FAILURE_COL_KEYWORDS = {
        "失效", "故障", "缺陷", "潜在失效", "失效模式", "失效原因", "故障模式", "失效后果",
        "failure", "defect", "fault", "potential", "failure_mode", "failure_effect"
    }
    CONTROL_COL_KEYWORDS = {
        "控制", "管控", "检查", "检验", "检测", "验证", "确认", "测量", "监控",
        "控制措施", "管控措施", "检验方法", "检测方法", "控制手段",
        "control", "inspection", "verify", "measure", "monitor", "control_method"
    }

    def _is_id_format(self, text: str) -> bool:
        text = str(text).strip()
        if re.match(r'^[A-Za-z]\d{3,}$', text):
            return True
        if re.match(r'^[A-Za-z]-\d+$', text):
            return True
        if re.match(r'^[A-Za-z]+-\d+-\d+$', text):
            return True
        if re.match(r'^8D\d{6,12}$', text):
            return True
        if re.match(r'^[A-Za-z0-9]+-\d{4}-\d{3,4}$', text):
            return True
        return False

    def _is_date_format(self, text: str) -> bool:
        if re.match(r'^\d{4}-\d{1,2}-\d{1,2}$', text):
            return True
        if re.match(r'^\d{4}/\d{1,2}/\d{1,2}$', text):
            return True
        if text == "NaT":
            return True
        return False

    def _is_meaningful(self, text: str) -> bool:
        if pd.isna(text):
            return False
        text = str(text).strip()
        if len(text) < 2 or len(text) > 120:
            return False
        if self._is_id_format(text):
            return False
        if self._is_date_format(text):
            return False
        if re.match(r'^[\d\s\.\-\(\)（）：:、，,；;]+$', text):
            return False
        has_chinese = bool(re.search(r'[\u4e00-\u9fff]', text))
        has_english = bool(re.search(r'[a-zA-Z]', text))
        if not (has_chinese or has_english):
            return False
        return True

    def _clean_text(self, text: str) -> str:
        text = str(text).strip()
        text = re.sub(r'\n.*$', '', text)
        text = re.sub(r'\s+', ' ', text)
        text = re.sub(r'^\d+[\.\)、]\s*', '', text)
        text = re.sub(r'^[（(]?[一二三四五六七八九十]+[）)]\s*', '', text)
        text = re.sub(r'^\[\d+\]\s*', '', text)
        text = re.sub(r'^No\.?\s*\d+\s*', '', text, flags=re.IGNORECASE)
        text = re.sub(r'^\s*(工序|特性|失效|控制|过程|步骤)[\d：:、，,；;\s]*', '', text)
        return text.strip()

    def _extract_chinese(self, text: str) -> str:
        text = str(text).strip()
        text = re.sub(r'\n.*$', '', text)
        chinese_parts = re.findall(r'[\u4e00-\u9fff][\u4e00-\u9fff\w\s]*', text)
        return ''.join(chinese_parts).strip() if chinese_parts else text

    def _find_column(self, columns: List[str], keywords: set) -> int:
        priority_idx = -1
        normal_idx = -1
        
        for i, col in enumerate(columns):
            col_str = str(col).lower()
            for kw in keywords:
                if kw.lower() in col_str:
                    if "名称" in str(col) or "name" in col_str or "项目" in str(col) or "process" in col_str:
                        priority_idx = i
                    else:
                        if normal_idx < 0:
                            normal_idx = i
                    break
        
        if priority_idx >= 0:
            return priority_idx
        return normal_idx

    def _score_header(self, header: str, keywords: set) -> int:
        header_lower = str(header).lower()
        score = 0
        for kw in keywords:
            kw_lower = kw.lower()
            if kw_lower in header_lower:
                score += 10
                if "名称" in header or "name" in header_lower or "项目" in header or "item" in header_lower:
                    score += 5
        return score

    def _extract_from_excel(self, file_path: str) -> Dict[str, List[str]]:
        result = {"processes": [], "characteristics": [], "failure_modes": [], "controls": []}
        
        raw_data = self._read_xlsx_raw(file_path)
        for sheet_name, data in raw_data.items():
            try:
                if len(data) < 2:
                    continue
                
                header_rows = self._detect_header_rows(data)
                if not header_rows:
                    continue
                
                merged_header = self._merge_header_rows(data, header_rows)
                data_start = max(header_rows) + 1
                rows = data[data_start:]
                
                self._extract_from_complex_rows(merged_header, rows, result)
                
            except Exception:
                continue
        
        return result

    def _detect_header_rows(self, data: List[List[Any]]) -> List[int]:
        if len(data) < 2:
            return []
        
        candidates = []
        for i, row in enumerate(data[:15]):
            row_text = ' '.join(str(c) for c in row if c)
            
            category_count = 0
            if any(kw in row_text for kw in self.PROCESS_COL_KEYWORDS):
                category_count += 1
            if any(kw in row_text for kw in self.CHARACTERISTIC_COL_KEYWORDS):
                category_count += 1
            if any(kw in row_text for kw in self.FAILURE_COL_KEYWORDS):
                category_count += 1
            if any(kw in row_text for kw in self.CONTROL_COL_KEYWORDS):
                category_count += 1
            
            header_kw_count = 0
            for kw in list(self.PROCESS_COL_KEYWORDS) + list(self.CHARACTERISTIC_COL_KEYWORDS) + list(self.FAILURE_COL_KEYWORDS) + list(self.CONTROL_COL_KEYWORDS):
                if kw in row_text:
                    header_kw_count += 1
            
            non_empty_count = len([c for c in row if c])
            
            is_likely_header = (category_count >= 2) or (header_kw_count >= 3 and non_empty_count >= 3)
            
            if is_likely_header:
                candidates.append((i, category_count, header_kw_count))
        
        if not candidates:
            return [0]
        
        candidates.sort(key=lambda x: (x[1], x[2]), reverse=True)
        best_start = candidates[0][0]
        
        header_rows = [best_start]
        for i in range(best_start + 1, min(best_start + 4, len(data))):
            row_text = ' '.join(str(c) for c in data[i] if c)
            is_subheader = any(kw in row_text for kw in ["项目", "级别", "level", "no.", "编号", "products", "characteristics"])
            mostly_empty = len([c for c in data[i] if c]) <= 4
            if is_subheader or mostly_empty:
                header_rows.append(i)
            else:
                break
        
        return header_rows

    def _merge_header_rows(self, data: List[List[Any]], header_rows: List[int]) -> List[str]:
        if not header_rows:
            return []
        
        max_len = 0
        for i in header_rows:
            max_len = max(max_len, len(data[i]))
        
        merged = [''] * max_len
        for i in header_rows:
            row = data[i]
            for j in range(min(len(row), max_len)):
                cell = str(row[j]).strip()
                if cell:
                    if merged[j]:
                        merged[j] = merged[j] + ' ' + cell
                    else:
                        merged[j] = cell
        
        return merged

    def _extract_from_complex_rows(self, header: List[str], rows: List[List[Any]], result):
        n_cols = len(header)
        
        header_texts = [str(h).lower() for h in header]
        
        process_name_idx = -1
        process_step_name_idx = -1
        char_indices = []
        failure_indices = []
        control_indices = []
        
        for i, h in enumerate(header):
            h_str = str(h)
            h_lower = h_str.lower()
            
            score_process = self._score_header(h_str, self.PROCESS_COL_KEYWORDS)
            score_char = self._score_header(h_str, self.CHARACTERISTIC_COL_KEYWORDS)
            score_failure = self._score_header(h_str, self.FAILURE_COL_KEYWORDS)
            score_control = self._score_header(h_str, self.CONTROL_COL_KEYWORDS)
            
            if score_process > 0:
                if "4m" in h_lower or "工作要素" in h_str:
                    continue
                if "步骤" in h_str or "step" in h_lower:
                    if "名称" in h_str or "name" in h_lower or "要素" in h_str or "关注" in h_str:
                        if process_step_name_idx < 0 or score_process > self._score_header(header[process_step_name_idx], self.PROCESS_COL_KEYWORDS):
                            process_step_name_idx = i
                elif "名称" in h_str or "name" in h_lower or "过程" in h_str or "process" in h_lower or "工位" in h_str:
                    if process_name_idx < 0 or score_process > self._score_header(header[process_name_idx], self.PROCESS_COL_KEYWORDS):
                        process_name_idx = i
            
            if score_char > 0 and not any(bad in h_str for bad in ["级别", "level", "频度", "严重度", "探测度", "公差", "tolerance"]):
                char_indices.append(i)
            if score_failure > 0 and "影响" not in h_str and "effect" not in h_lower and "严重度" not in h_str and "severity" not in h_lower:
                failure_indices.append(i)
            if score_control > 0 and not any(bad in h_str for bad in ["执行人员", "responsible", "目标完成时间", "状态", "人员", "operator"]):
                control_indices.append(i)
        
        if process_step_name_idx >= 0:
            process_idx = process_step_name_idx
        elif process_name_idx >= 0:
            process_idx = process_name_idx
        else:
            process_idx = -1
        
        if process_idx >= 0:
            process_idx = self._select_best_process_column(process_idx, rows, n_cols)
        
        last_process = ""
        for row in rows:
            if len(row) < n_cols:
                row = row + [''] * (n_cols - len(row))
            
            if process_idx >= 0 and process_idx < len(row):
                val = row[process_idx]
                if not val or str(val).strip() == '':
                    val = last_process
                else:
                    last_process = val
                
                if process_idx + 1 < len(row) and val and self._looks_like_code(val):
                    next_val = row[process_idx + 1]
                    if next_val and not self._looks_like_code(next_val):
                        val = next_val
                        last_process = val
                
                cleaned = self._clean_text(val)
                chinese = self._extract_chinese(cleaned)
                target = chinese if chinese else cleaned
                if self._is_meaningful(target) and target not in result["processes"]:
                    if not any(target.lower() == ht for ht in header_texts):
                        result["processes"].append(target)
            
            for idx in char_indices:
                if idx < len(row):
                    val = row[idx]
                    cleaned = self._clean_text(val)
                    chinese = self._extract_chinese(cleaned)
                    target = chinese if chinese else cleaned
                    if self._is_meaningful(target) and target not in result["characteristics"]:
                        if target not in ["SC", "CC", "AA", "BB"]:
                            result["characteristics"].append(target)
            
            for idx in failure_indices:
                if idx < len(row):
                    val = row[idx]
                    cleaned = self._clean_text(val)
                    chinese = self._extract_chinese(cleaned)
                    target = chinese if chinese else cleaned
                    if self._is_meaningful(target) and target not in result["failure_modes"]:
                        if not target.endswith("符合要求") and not target.endswith("meets the requirements"):
                            result["failure_modes"].append(target)
            
            for idx in control_indices:
                if idx < len(row):
                    val = row[idx]
                    cleaned = self._clean_text(val)
                    chinese = self._extract_chinese(cleaned)
                    target = chinese if chinese else cleaned
                    if self._is_meaningful(target) and target not in result["controls"]:
                        result["controls"].append(target)

    def _looks_like_code(self, text: str) -> bool:
        text = str(text).strip()
        if re.match(r'^[A-Za-z]\d{3,}$', text):
            return True
        if re.match(r'^\d+[\.\)、]?$', text):
            return True
        return False

    def _select_best_process_column(self, process_idx: int, rows: List[List[Any]], n_cols: int) -> int:
        candidates = [process_idx]
        if process_idx + 1 < n_cols:
            candidates.append(process_idx + 1)
        if process_idx - 1 >= 0:
            candidates.append(process_idx - 1)
        
        best_idx = process_idx
        best_score = -1
        for idx in candidates:
            chinese_count = 0
            code_count = 0
            empty_count = 0
            total = 0
            for row in rows:
                if idx < len(row):
                    val = str(row[idx]).strip()
                    if val:
                        total += 1
                        if re.search(r'[\u4e00-\u9fff]', val):
                            chinese_count += 1
                        if self._looks_like_code(val):
                            code_count += 1
                    else:
                        empty_count += 1
            
            if total == 0:
                continue
            score = chinese_count - code_count * 2
            if score > best_score:
                best_score = score
                best_idx = idx
        
        return best_idx

    def _read_xlsx_raw(self, file_path: str) -> Dict[str, List[List[Any]]]:
        sheets = {}
        ns = {'main': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        
        with zipfile.ZipFile(file_path, 'r') as z:
            shared_strings = []
            if 'xl/sharedStrings.xml' in z.namelist():
                sst_xml = z.read('xl/sharedStrings.xml')
                sst_root = ET.fromstring(sst_xml)
                for si in sst_root.findall('.//main:si', ns):
                    text_parts = []
                    for t in si.findall('.//main:t', ns):
                        text_parts.append(t.text or '')
                    shared_strings.append(''.join(text_parts))
            
            workbook_xml = z.read('xl/workbook.xml')
            root = ET.fromstring(workbook_xml)
            sheet_nodes = root.findall('.//main:sheet', ns)
            
            for sheet_node in sheet_nodes:
                sheet_name = sheet_node.get('name')
                sheet_id = sheet_node.get('sheetId')
                sheet_file = f'xl/worksheets/sheet{sheet_id}.xml'
                
                if sheet_file not in z.namelist():
                    continue
                
                sheet_xml = z.read(sheet_file)
                sheet_root = ET.fromstring(sheet_xml)
                
                merges = []
                merge_cells = sheet_root.find('.//main:mergeCells', ns)
                if merge_cells is not None:
                    for mc in merge_cells.findall('.//main:mergeCell', ns):
                        ref = mc.get('ref')
                        if ref:
                            merges.append(ref)
                
                row_dict = {}
                max_col_idx = 0
                for row in sheet_root.findall('.//main:row', ns):
                    row_idx = int(row.get('r'))
                    cells = {}
                    for cell in row.findall('.//main:c', ns):
                        ref = cell.get('r')
                        col_letter = re.match(r'([A-Z]+)', ref).group(1) if ref else ''
                        col_idx = self._col_to_num(col_letter)
                        max_col_idx = max(max_col_idx, col_idx)
                        cell_type = cell.get('t')
                        value_node = cell.find('.//main:v', ns)
                        text_nodes = cell.findall('.//main:t', ns)
                        
                        if text_nodes:
                            value = ''.join(t.text or '' for t in text_nodes)
                        elif value_node is not None and value_node.text:
                            raw_value = value_node.text
                            if cell_type == 's' and shared_strings:
                                try:
                                    value = shared_strings[int(raw_value)]
                                except (ValueError, IndexError):
                                    value = raw_value
                            else:
                                value = raw_value
                        else:
                            value = ''
                        
                        cells[col_idx] = value
                    
                    row_dict[row_idx] = cells
                
                header_row = self._estimate_header_row(row_dict, max_col_idx)
                
                for merge_ref in merges:
                    start, end = merge_ref.split(':')
                    start_col = re.match(r'([A-Z]+)', start).group(1)
                    start_row = int(re.search(r'(\d+)', start).group(1))
                    end_col = re.match(r'([A-Z]+)', end).group(1)
                    end_row = int(re.search(r'(\d+)', end).group(1))
                    start_c = self._col_to_num(start_col)
                    end_c = self._col_to_num(end_col)
                    
                    value = ''
                    if start_row in row_dict and start_c in row_dict[start_row]:
                        value = row_dict[start_row][start_c]
                    
                    for r in range(start_row, end_row + 1):
                        if r not in row_dict:
                            row_dict[r] = {}
                        for c in range(start_c, end_c + 1):
                            if c not in row_dict[r] or row_dict[r][c] == '':
                                if start_row <= header_row or r > header_row:
                                    row_dict[r][c] = value
                
                if row_dict:
                    max_row = max(row_dict.keys())
                    rows = []
                    for r in range(1, max_row + 1):
                        cells = row_dict.get(r, {})
                        row_data = [''] * (max_col_idx + 1)
                        for c, value in cells.items():
                            if c < len(row_data):
                                row_data[c] = value
                        rows.append(row_data)
                    
                    sheets[sheet_name] = rows
        
        return sheets

    def _estimate_header_row(self, row_dict, max_col_idx):
        best_row = 0
        best_score = -1
        for row_idx, cells in row_dict.items():
            if row_idx > 15:
                continue
            row_text = ' '.join(str(cells.get(c, '')) for c in range(max_col_idx + 1))
            score = 0
            if any(kw in row_text for kw in self.PROCESS_COL_KEYWORDS):
                score += 10
            if any(kw in row_text for kw in self.CHARACTERISTIC_COL_KEYWORDS):
                score += 10
            if any(kw in row_text for kw in self.FAILURE_COL_KEYWORDS):
                score += 10
            if any(kw in row_text for kw in self.CONTROL_COL_KEYWORDS):
                score += 10
            non_empty = len([c for c in cells.values() if c])
            score += min(non_empty, 5)
            if score > best_score:
                best_score = score
                best_row = row_idx
        return best_row

    def _col_to_num(self, col: str) -> int:
        num = 0
        for c in col:
            num = num * 26 + (ord(c) - ord('A') + 1)
        return num - 1

    def _fallback_extract(self, df, result):
        text = df.to_string(index=False)
        lines = text.strip().split("\n")

        for line in lines:
            line_stripped = line.strip()
            if not line_stripped:
                continue

            if len(result["processes"]) < 30:
                process_patterns = [
                    r'(?:工序|过程|步骤|工位|操作|流程)\s*[：:、，]?\s*([^\d：:、，;；\s]{2,30})',
                    r'^\s*\d+\s*[\.\)、]\s*([^\d：:、，;；\s]{2,30})',
                    r'(?:加工|装配|焊接|冲压|注塑|检验|测试|包装|组装|热处理|涂装)\s*[^\d\s]{0,20}'
                ]
                for pattern in process_patterns:
                    matches = re.findall(pattern, line_stripped)
                    for match in matches:
                        cleaned = self._clean_text(match)
                        if self._is_meaningful(cleaned) and cleaned not in result["processes"]:
                            result["processes"].append(cleaned)
                            break

            if len(result["failure_modes"]) < 30:
                failure_patterns = [
                    r'(?:失效|故障|缺陷|损坏|不良|异常)\s*[：:、，]?\s*([^\d：:、，;；\s]{2,40})',
                    r'(?:可能导致|引起|造成)\s*([^\d：:、，;；\s]{2,40})',
                    r'(?:出现|发生|产生)\s*(?:缺陷|故障|失效|不良)\s*[：:、，]?\s*([^\d：:、，;；\s]{2,40})'
                ]
                for pattern in failure_patterns:
                    matches = re.findall(pattern, line_stripped)
                    for match in matches:
                        cleaned = self._clean_text(match)
                        if self._is_meaningful(cleaned) and cleaned not in result["failure_modes"]:
                            result["failure_modes"].append(cleaned)
                            break

            if len(result["controls"]) < 30:
                control_patterns = [
                    r'(?:控制|检查|检验|检测|验证|确认|监控)\s*[：:、，]?\s*([^\d：:、，;；\s]{2,40})',
                    r'(?:措施|方法|手段)\s*[：:、，]?\s*([^\d：:、，;；\s]{2,40})',
                    r'(?:采用|使用|实施)\s*(?:控制|检查|检验|检测)\s*[^\d\s]{0,10}\s*([^\d：:、，;；\s]{2,40})'
                ]
                for pattern in control_patterns:
                    matches = re.findall(pattern, line_stripped)
                    for match in matches:
                        cleaned = self._clean_text(match)
                        if self._is_meaningful(cleaned) and cleaned not in result["controls"]:
                            result["controls"].append(cleaned)
                            break

    def _extract_from_text(self, text: str) -> Dict[str, List[str]]:
        result = {"processes": [], "characteristics": [], "failure_modes": [], "controls": []}
        lines = text.strip().split("\n")

        for line in lines:
            line_stripped = line.strip()
            if not line_stripped:
                continue

            if len(result["processes"]) < 30:
                process_patterns = [
                    r'(?:工序|过程|步骤|工位|操作|流程)\s*[：:、，]?\s*([^\d：:、，;；\s]{2,30})',
                    r'^\s*\d+\s*[\.\)、]\s*([^\d：:、，;；\s]{2,30})'
                ]
                for pattern in process_patterns:
                    matches = re.findall(pattern, line_stripped)
                    for match in matches:
                        cleaned = self._clean_text(match)
                        if self._is_meaningful(cleaned) and cleaned not in result["processes"]:
                            result["processes"].append(cleaned)

            if len(result["characteristics"]) < 30:
                char_patterns = [
                    r'(?:特性|尺寸|规格|参数)\s*[：:、，]?\s*([^\d：:、，;；\s]{2,30})'
                ]
                for pattern in char_patterns:
                    matches = re.findall(pattern, line_stripped)
                    for match in matches:
                        cleaned = self._clean_text(match)
                        if self._is_meaningful(cleaned) and cleaned not in result["characteristics"]:
                            result["characteristics"].append(cleaned)

            if len(result["failure_modes"]) < 30:
                failure_patterns = [
                    r'(?:失效|故障|缺陷)\s*[：:、，]?\s*([^\d：:、，;；\s]{2,40})'
                ]
                for pattern in failure_patterns:
                    matches = re.findall(pattern, line_stripped)
                    for match in matches:
                        cleaned = self._clean_text(match)
                        if self._is_meaningful(cleaned) and cleaned not in result["failure_modes"]:
                            result["failure_modes"].append(cleaned)

            if len(result["controls"]) < 30:
                control_patterns = [
                    r'(?:控制|检查|检验|检测|验证)\s*[：:、，]?\s*([^\d：:、，;；\s]{2,40})'
                ]
                for pattern in control_patterns:
                    matches = re.findall(pattern, line_stripped)
                    for match in matches:
                        cleaned = self._clean_text(match)
                        if self._is_meaningful(cleaned) and cleaned not in result["controls"]:
                            result["controls"].append(cleaned)

        return result

    def parse_document(self, file_path: str, doc_type: str) -> Dict[str, Any]:
        ext = file_path.lower().split(".")[-1]

        if ext in ("xlsx", "xls"):
            data = self._extract_from_excel(file_path)
        elif ext == "docx":
            data = self._extract_from_docx(file_path)
        else:
            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    text = f.read()
                data = self._extract_from_text(text)
            except Exception:
                data = {"processes": [], "characteristics": [], "failure_modes": [], "controls": []}

        result = {
            "raw_text": "",
            "lines": [],
            "processes": data["processes"],
            "characteristics": data["characteristics"],
            "failure_modes": data["failure_modes"],
            "controls": data["controls"]
        }

        if doc_type == "fmea":
            if not data["failure_modes"] and data["processes"]:
                result["failure_modes"] = data["processes"].copy()
        elif doc_type == "control_plan":
            if not data["controls"] and data["processes"]:
                result["controls"] = data["processes"].copy()

        return result

    def _extract_from_docx(self, file_path: str) -> Dict[str, List[str]]:
        result = {"processes": [], "characteristics": [], "failure_modes": [], "controls": []}
        try:
            doc = docx.Document(file_path)

            for table in doc.tables:
                try:
                    rows = []
                    for row in table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        rows.append(cells)

                    if len(rows) < 2:
                        continue

                    headers = rows[0]
                    process_col_idx = -1
                    char_col_idx = -1
                    failure_col_idx = -1
                    control_col_idx = -1

                    for i, header in enumerate(headers):
                        header_lower = str(header).lower()
                        for kw in self.PROCESS_COL_KEYWORDS:
                            if kw.lower() in header_lower:
                                if ("名称" in header or "name" in header_lower):
                                    process_col_idx = i
                                    break
                        if process_col_idx < 0:
                            for kw in self.PROCESS_COL_KEYWORDS:
                                if kw.lower() in header_lower:
                                    process_col_idx = i
                                    break

                        for kw in self.CHARACTERISTIC_COL_KEYWORDS:
                            if kw.lower() in header_lower:
                                char_col_idx = i
                                break
                        for kw in self.FAILURE_COL_KEYWORDS:
                            if kw.lower() in header_lower:
                                failure_col_idx = i
                                break
                        for kw in self.CONTROL_COL_KEYWORDS:
                            if kw.lower() in header_lower:
                                control_col_idx = i
                                break

                    for row in rows[1:]:
                        if process_col_idx >= 0 and process_col_idx < len(row):
                            val = row[process_col_idx]
                            cleaned = self._clean_text(val)
                            if self._is_meaningful(cleaned) and cleaned not in result["processes"]:
                                result["processes"].append(cleaned)

                        if char_col_idx >= 0 and char_col_idx < len(row):
                            val = row[char_col_idx]
                            cleaned = self._clean_text(val)
                            if self._is_meaningful(cleaned) and cleaned not in result["characteristics"]:
                                result["characteristics"].append(cleaned)

                        if failure_col_idx >= 0 and failure_col_idx < len(row):
                            val = row[failure_col_idx]
                            cleaned = self._clean_text(val)
                            if self._is_meaningful(cleaned) and cleaned not in result["failure_modes"]:
                                result["failure_modes"].append(cleaned)

                        if control_col_idx >= 0 and control_col_idx < len(row):
                            val = row[control_col_idx]
                            cleaned = self._clean_text(val)
                            if self._is_meaningful(cleaned) and cleaned not in result["controls"]:
                                result["controls"].append(cleaned)

                except Exception:
                    continue

            all_paragraphs = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            text_result = self._extract_from_text(all_paragraphs)

            for key in ["processes", "characteristics", "failure_modes", "controls"]:
                for item in text_result[key]:
                    if item not in result[key]:
                        result[key].append(item)

        except Exception:
            pass

        return result

    def _simple_match(self, source_item: str, target_items: List[str], min_overlap: int = 2) -> Tuple[bool, str]:
        if not source_item or not target_items:
            return False, ""

        source_clean = source_item.strip()

        for target in target_items:
            target_clean = target.strip()
            if not target_clean:
                continue

            if source_clean == target_clean:
                return True, target

            if source_clean in target_clean or target_clean in source_clean:
                if len(source_clean) >= 2 and len(target_clean) >= 2:
                    shorter = min(len(source_clean), len(target_clean))
                    longer = max(len(source_clean), len(target_clean))
                    if shorter / longer >= 0.5:
                        return True, target

            source_chars = set(c for c in source_clean if re.match(r'[\u4e00-\u9fff]', c))
            target_chars = set(c for c in target_clean if re.match(r'[\u4e00-\u9fff]', c))
            if source_chars and target_chars:
                overlap = len(source_chars & target_chars)
                shorter_len = min(len(source_chars), len(target_chars))
                if shorter_len > 0 and overlap / shorter_len >= 0.6:
                    return True, target

        return False, ""

    def validate(self, pfd_data: Dict[str, Any], fmea_data: Dict[str, Any], control_plan_data: Dict[str, Any] = None) -> Dict[str, Any]:
        results = {
            "summary": {
                "pfd_process_count": len(pfd_data["processes"]),
                "pfd_characteristic_count": len(pfd_data["characteristics"]),
                "fmea_failure_count": len(fmea_data["failure_modes"]),
                "fmea_process_count": len(fmea_data["processes"]),
                "control_plan_control_count": len(control_plan_data["controls"]) if control_plan_data else 0
            },
            "checks": []
        }

        checks = []

        checks.append({
            "category": "工序-失效匹配",
            "description": "检查PFD中的工序是否在FMEA中有对应失效分析",
            "passed": [],
            "missing": []
        })

        for process in pfd_data["processes"]:
            matched, found = self._simple_match(process, fmea_data["failure_modes"])
            if matched:
                checks[0]["passed"].append({
                    "pfd_item": process,
                    "fmea_match": found
                })
            else:
                matched_process, found_process = self._simple_match(process, fmea_data["processes"])
                if matched_process:
                    checks[0]["passed"].append({
                        "pfd_item": process,
                        "fmea_match": f"FMEA中存在同名工序: {found_process}"
                    })
                else:
                    checks[0]["missing"].append({
                        "pfd_item": process,
                        "reason": "FMEA中未找到对应失效条目或同名工序"
                    })

        checks.append({
            "category": "特性-失效匹配",
            "description": "检查PFD中的特性是否在FMEA中有对应失效分析",
            "passed": [],
            "missing": []
        })

        for char in pfd_data["characteristics"]:
            matched, found = self._simple_match(char, fmea_data["failure_modes"])
            if matched:
                checks[1]["passed"].append({
                    "pfd_item": char,
                    "fmea_match": found
                })
            else:
                matched_char, found_char = self._simple_match(char, fmea_data["characteristics"])
                if matched_char:
                    checks[1]["passed"].append({
                        "pfd_item": char,
                        "fmea_match": f"FMEA中存在同名特性: {found_char}"
                    })
                else:
                    checks[1]["missing"].append({
                        "pfd_item": char,
                        "reason": "FMEA中未找到对应失效条目或同名特性"
                    })

        if control_plan_data:
            checks.append({
                "category": "FMEA失效-控制计划匹配",
                "description": "检查FMEA中的失效模式是否在控制计划中有对应管控措施",
                "passed": [],
                "missing": []
            })

            for failure in fmea_data["failure_modes"]:
                matched, found = self._simple_match(failure, control_plan_data["controls"])
                if matched:
                    checks[2]["passed"].append({
                        "fmea_item": failure,
                        "control_plan_match": found
                    })
                else:
                    checks[2]["missing"].append({
                        "fmea_item": failure,
                        "reason": "控制计划中未找到对应管控措施"
                    })

            checks.append({
                "category": "工序-控制计划匹配",
                "description": "检查PFD中的工序是否在控制计划中有对应管控措施",
                "passed": [],
                "missing": []
            })

            for process in pfd_data["processes"]:
                matched, found = self._simple_match(process, control_plan_data["controls"])
                if matched:
                    checks[3]["passed"].append({
                        "pfd_item": process,
                        "control_plan_match": found
                    })
                else:
                    matched_process, found_process = self._simple_match(process, control_plan_data["processes"])
                    if matched_process:
                        checks[3]["passed"].append({
                            "pfd_item": process,
                            "control_plan_match": f"控制计划中存在同名工序: {found_process}"
                        })
                    else:
                        checks[3]["missing"].append({
                            "pfd_item": process,
                            "reason": "控制计划中未找到对应管控措施或同名工序"
                        })

        results["checks"] = checks

        return results

    def format_report(self, validation_result: Dict[str, Any]) -> str:
        report_lines = []
        report_lines.append("=" * 60)
        report_lines.append("流程图-FMEA-控制计划 三位一体跨文档链路校验报告")
        report_lines.append("=" * 60)
        report_lines.append("")

        summary = validation_result["summary"]
        report_lines.append("【文档概览】")
        report_lines.append(f"  PFD流程图工序数: {summary['pfd_process_count']}")
        report_lines.append(f"  PFD流程图特性数: {summary['pfd_characteristic_count']}")
        report_lines.append(f"  FMEA失效模式数: {summary['fmea_failure_count']}")
        report_lines.append(f"  FMEA工序数: {summary['fmea_process_count']}")
        if summary.get("control_plan_control_count") is not None:
            report_lines.append(f"  控制计划管控措施数: {summary['control_plan_control_count']}")
        report_lines.append("")

        for check in validation_result["checks"]:
            report_lines.append("-" * 60)
            report_lines.append(f"【校验项】{check['category']}")
            report_lines.append(f"  {check['description']}")
            report_lines.append("")

            if check["passed"]:
                report_lines.append(f"  ✅ 已匹配 ({len(check['passed'])}项):")
                for item in check["passed"]:
                    if "pfd_item" in item:
                        match_info = item.get("fmea_match", item.get("control_plan_match", ""))
                        report_lines.append(f"    - {item['pfd_item']} -> {match_info}")
                    elif "fmea_item" in item:
                        report_lines.append(f"    - {item['fmea_item']} -> {item.get('control_plan_match', '')}")
                report_lines.append("")

            if check["missing"]:
                report_lines.append(f"  ⚠️  缺失提示 ({len(check['missing'])}项):")
                for item in check["missing"]:
                    if "pfd_item" in item:
                        report_lines.append(f"    - PFD: {item['pfd_item']}")
                        report_lines.append(f"      原因: {item['reason']}")
                    elif "fmea_item" in item:
                        report_lines.append(f"    - FMEA: {item['fmea_item']}")
                        report_lines.append(f"      原因: {item['reason']}")
                report_lines.append("")

        report_lines.append("=" * 60)
        report_lines.append("【备注】")
        report_lines.append("  本校验仅基于显性文本匹配，不涉及语义理解。")
        report_lines.append("  校验结果仅供人工参考，请结合专业知识判断。")
        report_lines.append("=" * 60)

        return "\n".join(report_lines)