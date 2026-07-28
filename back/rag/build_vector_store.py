import argparse
import json
import logging
import os
from pathlib import Path
from typing import Any, Iterable, List, Optional

from .embeddings import get_embedding_model
from .vector_store import chunk_text
from .sqlite_vector_store import SqliteVectorStore

logger = logging.getLogger(__name__)

SUPPORTED_SUFFIXES = {".txt", ".md", ".json", ".pdf", ".xlsx"}

# 默认的本地 bge 模型路径：项目根目录的 models/bge-small-zh-v1.5
_DEFAULT_LOCAL_MODEL_DIR = Path(__file__).resolve().parent.parent.parent / "models" / "bge-small-zh-v1.5"


def _resolve_local_model_path() -> Optional[str]:
    """检测本地 bge 模型路径：优先用 EMBEDDING_MODEL_PATH，否则探测 models/ 目录。"""
    env_path = os.getenv("EMBEDDING_MODEL_PATH")
    if env_path and Path(env_path).is_dir():
        return env_path
    if _DEFAULT_LOCAL_MODEL_DIR.is_dir():
        return str(_DEFAULT_LOCAL_MODEL_DIR)
    return None


def _build_embedder(backend: str, dimension: int):
    """根据 backend 构造嵌入器，失败时由调用方决定是否终止。"""
    if backend == "semantic":
        local_path = _resolve_local_model_path()
        if local_path:
            return get_embedding_model(prefer="semantic", semantic_model=local_path, fallback_dimension=dimension)
    return get_embedding_model(prefer=backend, fallback_dimension=dimension)


def build_vector_store(
    input_dir: str | Path,
    output_path: str | Path,
    chunk_size: int = 800,
    chunk_overlap: int = 120,
    embedding_dimension: int = 384,
    backend: str = "semantic",
) -> int:
    """构建 JSONL 向量库（保留兼容旧路径）。"""
    source_dir = Path(input_dir)
    target_path = Path(output_path)
    embedder = _build_embedder(backend, embedding_dimension)

    files = sorted(
        path for path in source_dir.rglob("*") if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES
    )
    target_path.parent.mkdir(parents=True, exist_ok=True)

    count = 0
    with target_path.open("w", encoding="utf-8") as handle:
        for doc_index, path in enumerate(files, start=1):
            text = _read_document(path)
            for chunk_index, chunk in enumerate(chunk_text(text, chunk_size, chunk_overlap), start=1):
                record = {
                    "id": f"doc_{doc_index:03d}_chunk_{chunk_index:04d}",
                    "text": chunk,
                    "embedding": embedder.embed(chunk),
                    "metadata": {
                        "source": path.name,
                        "path": str(path),
                        "chunk_index": chunk_index,
                        "category": "APQP",
                    },
                }
                handle.write(json.dumps(record, ensure_ascii=False) + "\n")
                count += 1
    return count


def build_vector_store_sqlite(
    input_dir: str | Path,
    db_path: str | Path,
    chunk_size: int = 800,
    chunk_overlap: int = 120,
    embedding_dimension: int = 384,
    clear_existing: bool = True,
    backend: str = "semantic",
) -> int:
    """构建 SQLite 向量库。返回总分片数。

    backend: "semantic"（默认，使用 bge-small-zh-v1.5）或 "hashing"（兜底）。
    """
    source_dir = Path(input_dir)
    embedder = _build_embedder(backend, embedding_dimension)
    logger.info("构建 SQLite 向量库 backend=%s, dim=%d", backend, embedder.dimension)
    store = SqliteVectorStore(db_path)

    if clear_existing:
        for doc in store.list_documents():
            store.delete_document(doc.source)

    files = sorted(
        path for path in source_dir.rglob("*") if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES
    )

    count = 0
    for idx, path in enumerate(files, start=1):
        print(f"  [{idx}/{len(files)}] 读取 {path.name} ...", flush=True)
        text = _read_document(path)
        if not text.strip():
            print(f"  [{idx}/{len(files)}] 跳过 {path.name} (空内容)", flush=True)
            continue
        added = store.add_document(
            source=path.name,
            text=text,
            embedding_fn=embedder.embed,
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            path=str(path),
            category="APQP",
        )
        count += added
        print(f"  [{idx}/{len(files)}] ✓ {path.name}: {added} chunks", flush=True)
    return count


def _read_document(path: Path) -> str:
    if path.suffix.lower() == ".pdf":
        return _read_pdf(path)
    if path.suffix.lower() == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        return "\n".join(_iter_json_text(payload))
    if path.suffix.lower() == ".xlsx":
        return _read_xlsx(path)
    return path.read_text(encoding="utf-8")


def _read_xlsx(path: Path) -> str:
    """将 Excel 文件渲染为易读文本，供 RAG 嵌入使用。

    支持两种常见布局：
    1. 标准表格型：第一行表头 + 多行数据（如 sample xlsx）
    2. 表单/报告型：字段名在左列，值在右列，不规则布局（如 8D 报告）

    策略：逐行提取非空值，不强行要求第一行是表头；
    用 " | " 连接同一行的值，空行作为自然分隔。
    """
    try:
        from openpyxl import load_workbook
        from openpyxl.worksheet.dimensions import SheetFormatProperties
    except ImportError as exc:
        raise RuntimeError(
            "读取 xlsx 需要 openpyxl，请安装：pip install openpyxl"
        ) from exc

    # 兼容高版本 Excel/WPS 生成的 xlsx（openpyxl 3.1.x 不支持 defaultColWidthPt）
    _orig_sfp_init = SheetFormatProperties.__init__
    if not hasattr(_read_xlsx, "_patched"):
        _known_sfp = {
            "baseColWidth", "customHeight", "defaultColWidth", "defaultRowHeight",
            "outlineLevelCol", "outlineLevelRow", "thickBottom", "thickTop",
            "zeroHeight",
        }
        def _patched_sfp_init(self, *args, **kwargs):
            filtered = {k: v for k, v in kwargs.items() if k in _known_sfp}
            return _orig_sfp_init(self, *args, **filtered)
        SheetFormatProperties.__init__ = _patched_sfp_init
        _read_xlsx._patched = True  # type: ignore[attr-defined]

    workbook = load_workbook(path, data_only=True, read_only=True)
    sheets_text = []

    for worksheet in workbook.worksheets:
        lines: List[str] = [f"[Sheet: {worksheet.title}]"]
        prev_empty = True

        for raw_row in worksheet.iter_rows(values_only=True):
            # 提取非空值
            values = [
                str(v).strip()
                for v in raw_row
                if v is not None and str(v).strip() not in {"", "None"}
            ]
            if not values:
                if not prev_empty:
                    lines.append("")  # 空行作为段落分隔
                    prev_empty = True
                continue

            prev_empty = False
            # 2-4 个值的行通常是 "字段 | 值" 或简单表格行
            # 多于 4 个值的是标准表格数据行
            lines.append(" | ".join(values))

        # 去掉尾部空行
        while lines and lines[-1] == "":
            lines.pop()

        if len(lines) > 1:  # 至少有一个标题 + 内容
            sheets_text.append("\n".join(lines))

    workbook.close()

    if not sheets_text:
        raise RuntimeError(f"xlsx 解析后无有效数据：{path}")

    return "\n\n".join(sheets_text)


def _read_pdf(path: Path) -> str:
    """读取 PDF 文本，优先 PyMuPDF（fitz），缺失时回退到 PyPDF2。"""
    try:
        import fitz
        pages = []
        with fitz.open(path) as doc:
            for page_number, page in enumerate(doc, start=1):
                text = page.get_text("text", sort=True).strip()
                if text:
                    pages.append(f"[Page {page_number}]\n{text}")
        if pages:
            return "\n\n".join(pages)
    except ImportError:
        pass

    # 兜底：PyPDF2
    try:
        import PyPDF2
    except ImportError as exc:
        raise RuntimeError(
            "读取 PDF 需要 PyMuPDF 或 PyPDF2，请安装："
            "pip install PyMuPDF  # 或 pip install PyPDF2"
        ) from exc
    reader = PyPDF2.PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        text = text.strip()
        if text:
            pages.append(f"[Page {i}]\n{text}")
    if not pages:
        raise RuntimeError(f"PDF 解析后无文本：{path}")
    return "\n\n".join(pages)


def _iter_json_text(value: Any) -> Iterable[str]:
    if isinstance(value, str):
        text = value.strip()
        if text:
            yield text
    elif isinstance(value, list):
        for item in value:
            yield from _iter_json_text(item)
    elif isinstance(value, dict):
        for item in value.values():
            yield from _iter_json_text(item)


def main(argv: List[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build a local vector store from cleaned knowledge docs.")
    parser.add_argument("--input-dir", default="knowledge_docs")
    parser.add_argument("--output", default="knowledge_base/vector_store.sqlite")
    parser.add_argument("--format", choices=["jsonl", "sqlite"], default="sqlite",
                        help="输出格式，默认 sqlite（推荐，支持文档删除/管理）")
    parser.add_argument("--chunk-size", type=int, default=800)
    parser.add_argument("--chunk-overlap", type=int, default=120)
    parser.add_argument("--dimension", type=int, default=384,
                        help="哈希模式的向量维度；语义模式固定 512")
    parser.add_argument("--backend", choices=["semantic", "hashing"], default="semantic",
                        help="嵌入后端，默认 semantic (bge-small-zh-v1.5)")
    args = parser.parse_args(argv)

    if args.format == "sqlite":
        count = build_vector_store_sqlite(
            input_dir=args.input_dir,
            db_path=args.output,
            chunk_size=args.chunk_size,
            chunk_overlap=args.chunk_overlap,
            embedding_dimension=args.dimension,
            backend=args.backend,
        )
        print(f"Built {count} vector records at {args.output} (sqlite, backend={args.backend})")
    else:
        count = build_vector_store(
            input_dir=args.input_dir,
            output_path=args.output,
            chunk_size=args.chunk_size,
            chunk_overlap=args.chunk_overlap,
            embedding_dimension=args.dimension,
            backend=args.backend,
        )
        print(f"Built {count} vector records at {args.output} (jsonl, backend={args.backend})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
